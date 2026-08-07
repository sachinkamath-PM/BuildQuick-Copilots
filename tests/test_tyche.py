import json
from io import BytesIO
from uuid import uuid4

from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.services.auth import Identity, issue_token


client = TestClient(app)


def identity_headers() -> dict[str, str]:
    suffix = str(uuid4())
    token = issue_token(Identity(user_id=f"tyche-user-{suffix}", workspace_id=f"tyche-space-{suffix}"))
    return {"Authorization": f"Bearer {token}"}


def parse_complete(response) -> dict:
    events = [json.loads(line) for line in response.text.splitlines() if line]
    return next(event["message"] for event in events if event["type"] == "complete")


def resume_docx() -> bytes:
    buffer = BytesIO()
    document = Document()
    document.add_heading("Avery Shah", level=1)
    document.add_paragraph("avery@example.com · Bengaluru")
    document.add_heading("EXPERIENCE", level=1)
    document.add_paragraph("Product Lead · Atlas Systems · 2022–Present")
    document.add_paragraph(
        "Led roadmap prioritisation with product, engineering, and sales.",
        style="List Bullet",
    )
    document.add_paragraph(
        "Improved activation by 18% using customer research and analytics.",
        style="List Bullet",
    )
    document.add_heading("SKILLS", level=1)
    document.add_paragraph("Product strategy · SQL · Experimentation")
    document.save(buffer)
    return buffer.getvalue()


def test_job_description_creates_deterministic_ats_and_evidence() -> None:
    headers = identity_headers()
    response = client.put(
        "/api/tyche/job-description",
        headers=headers,
        json={
            "title": "Senior Product Manager",
            "company": "Acme",
            "text": (
                "We need a Senior Product Manager to lead product strategy, manage the product "
                "roadmap, partner with engineering and design, and use customer research and analytics."
            ),
        },
    )
    assert response.status_code == 200
    workspace = response.json()["workspace"]
    assert workspace["ats"]["deterministic"] is True
    assert "engineering" in workspace["ats"]["matched_keywords"]
    assert "product strategy" in workspace["ats"]["missing_keywords"]
    assert any(item["kind"] == "user_claim" for item in workspace["evidence"])
    assert any(item["kind"] == "job_requirement" for item in workspace["evidence"])


def test_resume_docx_import_preserves_structure_and_exports_current_claims() -> None:
    headers = identity_headers()
    imported = client.post(
        "/api/tyche/resume/import",
        headers=headers,
        files={
            "file": (
                "Avery-Shah-Resume.docx",
                resume_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert imported.status_code == 200
    workspace = imported.json()["workspace"]
    assert workspace["resume"]["source_filename"] == "Avery-Shah-Resume.docx"
    assert workspace["resume"]["candidate_name"] == "Avery Shah"
    assert [item["kind"] for item in workspace["resume"]["blocks"]].count("heading") == 3
    assert workspace["resume"]["bullets"][0]["text"].startswith("Led roadmap")
    assert any(
        item["source_id"] == workspace["resume"]["bullets"][1]["id"]
        and "18%" in item["excerpt"]
        for item in workspace["evidence"]
    )

    exported = client.get("/api/tyche/resume/export.docx", headers=headers)
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    exported_document = Document(BytesIO(exported.content))
    exported_text = "\n".join(paragraph.text for paragraph in exported_document.paragraphs)
    assert "Avery Shah" in exported_text
    assert "Led roadmap prioritisation" in exported_text
    assert "Improved activation by 18%" in exported_text


def test_imported_resume_claim_remains_grounded_through_apply_and_undo() -> None:
    headers = identity_headers()
    workspace = client.post(
        "/api/tyche/resume/import",
        headers=headers,
        files={"file": ("resume.docx", resume_docx())},
    ).json()["workspace"]
    bullet = workspace["resume"]["bullets"][0]
    conversation = client.post(
        "/api/conversations", headers=headers, json={"product": "tyche"}
    ).json()
    response = client.post(
        f"/api/conversations/{conversation['id']}/messages",
        headers=headers,
        json={
            "content": "Make this more senior",
            "context": {
                "product": "tyche",
                "page": {"type": "resume", "id": workspace["resume"]["id"]},
                "selection": {"type": "resume_bullet", "ids": [bullet["id"]]},
                "filters": {"grounded_workspace": True},
                "snapshot_version": "client-value-is-ignored",
            },
        },
    )
    proposal = parse_complete(response)["proposals"][0]
    assert proposal["diff"]["original"] == bullet["text"]
    applied = client.post(
        f"/api/proposals/{proposal['id']}",
        headers={**headers, "Idempotency-Key": "imported-resume-apply"},
        json={"action": "apply", "source_version": proposal["source_version"]},
    ).json()
    edited = next(
        item for item in applied["workspace"]["resume"]["blocks"] if item["id"] == bullet["id"]
    )
    assert edited["text"] == proposal["diff"]["suggested"]
    undone = client.post(
        f"/api/tyche/changes/{applied['undo_change_id']}/undo", headers=headers
    ).json()["workspace"]
    restored = next(item for item in undone["resume"]["blocks"] if item["id"] == bullet["id"])
    assert restored["text"] == bullet["text"]


def test_job_description_txt_import_and_unsupported_file_rejection() -> None:
    headers = identity_headers()
    text = (
        "Lead product strategy and customer research. Partner with engineering, design, and sales "
        "to own roadmap prioritisation, analytics, and experimentation."
    )
    imported = client.post(
        "/api/tyche/job-description/import",
        headers=headers,
        data={"title": "Product Director", "company": "Acme"},
        files={"file": ("role.txt", text.encode(), "text/plain")},
    )
    assert imported.status_code == 200
    assert imported.json()["workspace"]["job_description"]["text"] == text
    rejected = client.post(
        "/api/tyche/resume/import",
        headers=headers,
        files={"file": ("resume.rtf", b"{\\rtf1 unsafe}", "application/rtf")},
    )
    assert rejected.status_code == 415


def test_grounded_chat_apply_and_undo_round_trip() -> None:
    headers = identity_headers()
    workspace = client.get("/api/tyche/workspace", headers=headers).json()["workspace"]
    conversation = client.post(
        "/api/conversations", headers=headers, json={"product": "tyche"}
    ).json()
    context = {
        "product": "tyche",
        "page": {"type": "resume", "id": workspace["resume"]["id"], "label": "Product résumé"},
        "selection": {
            "type": "resume_bullet",
            "ids": ["bullet-1"],
            "label": "Selected résumé bullet",
            "excerpt": "client text must be replaced by server context",
        },
        "filters": {"grounded_workspace": True},
        "authorised_resources": [],
        "snapshot_version": "client-version-must-not-be-trusted",
    }
    message_response = client.post(
        f"/api/conversations/{conversation['id']}/messages",
        headers=headers,
        json={"content": "Make this more senior", "context": context},
    )
    assert message_response.status_code == 200
    message = parse_complete(message_response)
    proposal = message["proposals"][0]
    assert proposal["diff"]["original"] == "Managed the product roadmap."
    assert proposal["source_version"] == workspace["resume"]["id"] + ":v1"
    assert proposal["payload"]["tyche_managed"] is True
    assert "commercial" not in proposal["diff"]["suggested"].lower()

    applied = client.post(
        f"/api/proposals/{proposal['id']}",
        headers={**headers, "Idempotency-Key": "grounded-apply"},
        json={"action": "apply", "source_version": proposal["source_version"]},
    )
    assert applied.status_code == 200
    applied_body = applied.json()
    assert applied_body["workspace"]["resume"]["version"] == 2
    assert applied_body["workspace"]["resume"]["bullets"][0]["text"] == proposal["diff"]["suggested"]

    undone = client.post(
        f"/api/tyche/changes/{applied_body['undo_change_id']}/undo", headers=headers
    )
    assert undone.status_code == 200
    undone_workspace = undone.json()["workspace"]
    assert undone_workspace["resume"]["version"] == 3
    assert undone_workspace["resume"]["bullets"][0]["text"] == "Managed the product roadmap."


def test_undo_rejects_when_a_later_change_exists() -> None:
    headers = identity_headers()
    workspace = client.get("/api/tyche/workspace", headers=headers).json()["workspace"]
    conversation = client.post(
        "/api/conversations", headers=headers, json={"product": "tyche"}
    ).json()

    def propose_and_apply(idempotency_key: str) -> dict:
        current = client.get("/api/tyche/workspace", headers=headers).json()["workspace"]
        context = {
            "product": "tyche",
            "page": {"type": "resume", "id": current["resume"]["id"]},
            "selection": {"type": "resume_bullet", "ids": ["bullet-1"]},
            "filters": {"grounded_workspace": True},
            "snapshot_version": "ignored",
        }
        response = client.post(
            f"/api/conversations/{conversation['id']}/messages",
            headers=headers,
            json={"content": "Improve this", "context": context},
        )
        proposal = parse_complete(response)["proposals"][0]
        return client.post(
            f"/api/proposals/{proposal['id']}",
            headers={**headers, "Idempotency-Key": idempotency_key},
            json={"action": "apply", "source_version": proposal["source_version"]},
        ).json()

    first = propose_and_apply("first-change")
    propose_and_apply("second-change")
    stale_undo = client.post(
        f"/api/tyche/changes/{first['undo_change_id']}/undo", headers=headers
    )
    assert stale_undo.status_code == 409


def test_chat_explains_deterministic_ats_without_inventing_requirements() -> None:
    headers = identity_headers()
    saved = client.put(
        "/api/tyche/job-description",
        headers=headers,
        json={
            "title": "Senior Product Manager",
            "company": "Acme",
            "text": (
                "Lead product strategy and customer research while partnering with engineering, "
                "design, and sales to manage a product roadmap and experimentation programme."
            ),
        },
    ).json()["workspace"]
    conversation = client.post(
        "/api/conversations", headers=headers, json={"product": "tyche"}
    ).json()
    response = client.post(
        f"/api/conversations/{conversation['id']}/messages",
        headers=headers,
        json={
            "content": "Why is my ATS score low?",
            "context": {
                "product": "tyche",
                "page": {"type": "resume", "id": "resume-1"},
                "selection": {"type": "resume_bullet", "ids": ["bullet-1"]},
                "filters": {"grounded_workspace": True},
                "snapshot_version": "untrusted-client-version",
            },
        },
    )
    message = parse_complete(response)
    assert f"{saved['ats']['score']}/100" in message["content"]
    assert "Only add a missing requirement" in message["content"]
    assert message["proposals"] == []
