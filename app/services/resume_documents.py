from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from fastapi import HTTPException
from pypdf import PdfReader

from app.domain.models import ResumeBlock, ResumeDocument


MAX_UPLOAD_BYTES = 8 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 60_000
MAX_PDF_PAGES = 15
MAX_DOCX_UNCOMPRESSED_BYTES = 24 * 1024 * 1024
MAX_DOCX_ENTRIES = 1_500
SUPPORTED_EXTENSIONS = {".docx", ".pdf", ".txt"}
_BULLET_PREFIX = re.compile(r"^\s*(?:[•●▪◦‣⁃\uf0b7]|[-–—*])\s+")


@dataclass(frozen=True)
class ImportedDocument:
    filename: str
    blocks: list[ResumeBlock]
    candidate_name: str | None
    text: str


def safe_filename(filename: str | None) -> str:
    cleaned = Path(filename or "document").name.strip()
    if not cleaned or cleaned in {".", ".."}:
        raise HTTPException(status_code=422, detail="A valid filename is required")
    return cleaned[:255]


def read_upload(data: bytes, filename: str | None) -> ImportedDocument:
    name = safe_filename(filename)
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise HTTPException(status_code=415, detail="Upload a DOCX, PDF, or TXT file")
    if not data:
        raise HTTPException(status_code=422, detail="The uploaded file is empty")
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="The uploaded file exceeds the 8 MB limit")

    if suffix == ".docx":
        raw_blocks = _read_docx(data)
    elif suffix == ".pdf":
        raw_blocks = _read_pdf(data)
    else:
        raw_blocks = _read_text(data)
    blocks = _normalise_blocks(raw_blocks)
    if not blocks:
        raise HTTPException(status_code=422, detail="No readable text was found in the document")
    text = "\n".join(block.text for block in blocks)
    candidate = _candidate_name(blocks)
    return ImportedDocument(filename=name, blocks=blocks, candidate_name=candidate, text=text)


def _read_docx(data: bytes) -> list[tuple[str, str, int]]:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_DOCX_ENTRIES:
                raise HTTPException(status_code=413, detail="The DOCX contains too many embedded files")
            if sum(item.file_size for item in entries) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise HTTPException(status_code=413, detail="The expanded DOCX exceeds the safety limit")
        document = Document(BytesIO(data))
    except HTTPException:
        raise
    except (zipfile.BadZipFile, ValueError, KeyError) as exc:
        raise HTTPException(status_code=422, detail="The DOCX file could not be read") from exc

    blocks: list[tuple[str, str, int]] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name if paragraph.style else "").lower()
        has_numbering = paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None
        if style_name.startswith("heading"):
            match = re.search(r"(\d+)", style_name)
            blocks.append(("heading", text, min(3, int(match.group(1))) if match else 1))
        elif has_numbering or "list" in style_name or _BULLET_PREFIX.match(text):
            blocks.append(("bullet", _BULLET_PREFIX.sub("", text), 0))
        else:
            blocks.append(("paragraph", text, 0))
    for table in document.tables:
        for row in table.rows:
            text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if text:
                blocks.append(("paragraph", text, 0))
    return blocks


def _read_pdf(data: bytes) -> list[tuple[str, str, int]]:
    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise HTTPException(status_code=422, detail="Password-protected PDFs are not supported")
            except Exception as exc:
                raise HTTPException(status_code=422, detail="Password-protected PDFs are not supported") from exc
        if len(reader.pages) > MAX_PDF_PAGES:
            raise HTTPException(status_code=413, detail=f"PDFs are limited to {MAX_PDF_PAGES} pages")
        lines: list[str] = []
        for page in reader.pages:
            lines.extend((page.extract_text() or "").splitlines())
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail="The PDF file could not be read") from exc
    return [_classify_plain_line(line) for line in lines if line.strip()]


def _read_text(data: bytes) -> list[tuple[str, str, int]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise HTTPException(status_code=422, detail="TXT files must use UTF-8 encoding") from exc
    return [_classify_plain_line(line) for line in text.splitlines() if line.strip()]


def _classify_plain_line(line: str) -> tuple[str, str, int]:
    text = line.strip()
    if _BULLET_PREFIX.match(text):
        return "bullet", _BULLET_PREFIX.sub("", text), 0
    letters = "".join(character for character in text if character.isalpha())
    if 2 <= len(text.split()) <= 6 and letters and letters.upper() == letters:
        return "heading", text, 1
    return "paragraph", text, 0


def _normalise_blocks(raw: list[tuple[str, str, int]]) -> list[ResumeBlock]:
    trimmed: list[tuple[str, str, int]] = []
    total = 0
    for kind, text, level in raw:
        text = re.sub(r"\s+", " ", text).strip()
        if not text:
            continue
        if total + len(text) > MAX_EXTRACTED_CHARACTERS:
            raise HTTPException(status_code=413, detail="The document contains too much text")
        total += len(text)
        trimmed.append((kind, text[:2_000], level))

    if not any(kind == "bullet" for kind, _, _ in trimmed):
        candidates = [index for index, (kind, text, _) in enumerate(trimmed) if kind == "paragraph" and len(text.split()) >= 8]
        for index in candidates[:30]:
            kind, text, level = trimmed[index]
            trimmed[index] = ("bullet", text, level)

    blocks: list[ResumeBlock] = []
    bullet_number = 0
    for index, (kind, text, level) in enumerate(trimmed, start=1):
        if kind == "bullet":
            bullet_number += 1
            block_id = f"bullet-{bullet_number}"
        else:
            block_id = f"block-{index}"
        blocks.append(ResumeBlock(id=block_id, kind=kind, text=text, level=level))
    return blocks


def _candidate_name(blocks: list[ResumeBlock]) -> str | None:
    for block in blocks[:3]:
        words = block.text.split()
        if block.kind != "bullet" and 2 <= len(words) <= 5 and not any(char.isdigit() for char in block.text):
            return block.text
    return None


def export_resume_docx(resume: ResumeDocument) -> bytes:
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(35, 45, 51)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    for style_name, size, before, after in (("Heading 1", 15, 16, 7), ("Heading 2", 12, 11, 5)):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(28, 93, 81)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    if "Resume Name" not in styles:
        name_style = styles.add_style("Resume Name", WD_STYLE_TYPE.PARAGRAPH)
    else:
        name_style = styles["Resume Name"]
    name_style.font.name = "Calibri"
    name_style.font.size = Pt(22)
    name_style.font.bold = True
    name_style.font.color.rgb = RGBColor(24, 33, 29)
    name_style.paragraph_format.space_after = Pt(3)

    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.title = resume.title

    if resume.candidate_name:
        name = document.add_paragraph(resume.candidate_name, style="Resume Name")
        name.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if resume.target_role:
            target = document.add_paragraph(resume.target_role)
            target.runs[0].font.color.rgb = RGBColor(96, 108, 103)
            target.paragraph_format.space_after = Pt(12)

    bullet_lookup = {bullet.id: bullet.text for bullet in resume.bullets}
    blocks = resume.blocks or _default_blocks(resume)
    for block in blocks:
        text = bullet_lookup.get(block.id, block.text) if block.kind == "bullet" else block.text
        if resume.candidate_name and text == resume.candidate_name:
            continue
        if block.kind == "heading":
            document.add_paragraph(text, style="Heading 1" if block.level <= 1 else "Heading 2")
        elif block.kind == "bullet":
            paragraph = document.add_paragraph(text, style="List Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.keep_together = True
        else:
            document.add_paragraph(text)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _default_blocks(resume: ResumeDocument) -> list[ResumeBlock]:
    blocks = [
        ResumeBlock(id="default-experience", kind="heading", text="EXPERIENCE", level=1),
        ResumeBlock(id="default-company", kind="heading", text=resume.company, level=2),
        ResumeBlock(id="default-role", kind="paragraph", text=f"{resume.role} · {resume.period}"),
    ]
    blocks.extend(ResumeBlock(id=bullet.id, kind="bullet", text=bullet.text) for bullet in resume.bullets)
    return blocks
